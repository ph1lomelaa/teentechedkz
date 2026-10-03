"""Read-only XLSX extraction. Imported values are never published automatically."""
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook
from pydantic import ValidationError
from app.schemas.activity import ActivityCreate

MAX_BYTES = 5 * 1024 * 1024
HEADERS = ['№', 'Группа', 'Тип активности', 'Наименование', 'Организатор', 'Страна', 'Даты / дедлайн', 'Тема', 'Eligibility', 'Задание', 'Link']


def identity(title, organizer, url):
    def clean(value):
        return ' '.join((value or '').casefold().split())
    if url:
        parts = urlsplit(url)
        normalized = urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path.rstrip('/'), parts.query, ''))
        value = 'url:' + normalized
    else:
        value = 'name:' + clean(title) + '|' + clean(organizer)
    return hashlib.sha256(value.encode()).hexdigest()


def preview_xlsx(content, filename):
    if len(content) > MAX_BYTES:
        raise ValueError('Файл должен быть не больше 5 МБ')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError('Слишком большой распакованный файл')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError('Не удалось прочитать XLSX') from exc
    rows = []
    seen = set()
    skipped = 0
    try:
        sheet = book.worksheets[0]
        header = [str(c.value or '').strip() for c in next(sheet.iter_rows(min_row=1, max_row=1, max_col=11))]
        if header != HEADERS:
            raise ValueError('Колонки не совпадают с шаблоном базы активностей')
        if sheet.max_row and sheet.max_row > 10000:
            raise ValueError('Допускается не больше 10 000 строк')
        for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=11), start=2):
            if number > 10000:
                raise ValueError('Допускается не больше 10 000 строк')
            raw = {}
            warnings = []
            for key, cell in zip(HEADERS, cells):
                value = cell.value
                if cell.data_type == 'f':
                    warnings.append(f'Формула в колонке «{key}» не импортирована')
                    value = None
                if isinstance(value, (date, datetime)):
                    value = value.isoformat()
                raw[key] = str(value).strip() if value is not None else ''
            if not any(raw[k] for k in HEADERS[1:]):
                if raw['№']:
                    skipped += 1
                continue
            title = raw['Наименование']
            url = raw['Link'] or None
            errors = []
            if not title or re.match(r'^https?://', title, re.I):
                errors.append('Нужно указать название активности вместо пустого поля или ссылки')
            try:
                activity = ActivityCreate(title=title, direction=raw['Группа'] or None,
                    kind=raw['Тип активности'] or None, organizer=raw['Организатор'] or None,
                    country=raw['Страна'] or None, description=raw['Тема'] or None,
                    eligibility=raw['Eligibility'] or None, assignment=raw['Задание'] or None,
                    source_url=url).model_dump(mode='json')
            except ValidationError:
                errors.append('Проверьте название, длину полей и ссылку http/https')
                activity = None
            if not raw['Организатор']:
                warnings.append('Организатор не указан')
            if not url:
                warnings.append('Нет ссылки организатора')
            if raw['Даты / дедлайн']:
                warnings.append('Сроки сохранены текстом; точный дедлайн нужно проверить')
            key = identity(title, raw['Организатор'], activity['source_url'] if activity else url)
            if key in seen:
                errors.append('Дубликат в файле')
            seen.add(key)
            rows.append(dict(row=number, key=key, activity=activity, dates_text=raw['Даты / дедлайн'] or None,
                warnings=warnings, errors=errors, source=dict(filename=filename, sheet=sheet.title, row=number, raw=raw)))
        return dict(rows=rows, skipped=skipped)
    finally:
        book.close()
