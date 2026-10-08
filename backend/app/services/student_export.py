"""Read-only reconciliation rows for configurable student exports."""
from __future__ import annotations

import csv
import io
from collections import defaultdict

from openpyxl import Workbook

from app.models.contract import Contract
from app.models.application import Application
from app.models.mentor_assignment import MentorAssignment, MentorRole
from app.models.notion_snapshot import NotionMatchStatus, NotionSnapshot
from app.models.student import Student
from app.models.telegram_chat import TelegramChat, TelegramChatStatus
from app.models.telegram_chat_session import TelegramChatSession, TelegramSessionStatus
from app.models.user import User
from migration.transformers.normalize import parse_pipeline_status


COLUMNS = {
    "full_name": "Студент",
    "phone": "Телефон",
    "intake_year": "Год набора",
    "degree_level": "Уровень",
    "city": "Город",
    "email": "Email",
    "countries": "Страны",
    "programs": "Программы",
    "lead_mentor": "Lead-Mentor",
    "notion_status": "Статус Notion",
    "crm_status": "Статус CRM",
    "status_match": "Статусы совпадают",
    "notion_link": "Связь с Notion",
    "notion_url": "Ссылка на Notion",
    "portal_access": "Доступ к платформе",
    "active_chat": "Активный чат",
    "mzk": "МЗК",
    "mentors": "Менторы",
}
DEFAULT_COLUMNS = ("full_name", "phone", "notion_status", "crm_status", "portal_access", "active_chat", "mzk", "mentors")


def notion_status(snapshot: NotionSnapshot | None) -> tuple[str, str | None]:
    raw = (snapshot.normalized_data or {}).get("payment_status_raw") if snapshot else None
    if not raw:
        return "", None
    try:
        return str(raw), parse_pipeline_status(str(raw))
    except (ValueError, KeyError):
        return str(raw), None


def make_rows(
    students: list[Student], contracts: list[Contract], snapshots: list[NotionSnapshot],
    users: list[User], assignments: list[MentorAssignment],
    sessions: list[tuple[TelegramChatSession, TelegramChat]],
    dataset: str,
    applications: list[Application] | None = None,
    require_portal_access: bool = False,
    require_chat: bool = False,
    require_mzk: bool = False,
    require_mentor: bool = False,
    student_ids: set | None = None,
    portal_access: str = "any",
    active_chat: str = "any",
    has_mzk: str = "any",
    has_mentor: str = "any",
    notion_link: str = "any",
    notion_snapshot_ids: set | None = None,
    field_filters: dict[str, set[str]] | None = None,
    only_active_notion: bool = False,
    only_status_mismatch: bool = False,
    search: str = "",
) -> list[dict[str, str]]:
    user_by_id = {u.id: u for u in users}
    contracts_by_student: dict = {}
    for contract in sorted(contracts, key=lambda c: (c.created_at.isoformat() if c.created_at else "", str(c.id))):
        contracts_by_student[contract.student_id] = contract
    countries_by_student: dict = defaultdict(set)
    programs_by_student: dict = defaultdict(set)
    leads_by_student: dict = defaultdict(set)
    for application in applications or []:
        if application.country:
            countries_by_student[application.student_id].add(application.country)
        if application.program:
            programs_by_student[application.student_id].add(application.program)
        if application.lead_mentor_id and application.lead_mentor_id in user_by_id:
            leads_by_student[application.student_id].add(user_by_id[application.lead_mentor_id].name)
    snapshots_by_student: dict = {}
    for snapshot in sorted(snapshots, key=lambda s: (s.synced_at.isoformat() if s.synced_at else "", str(s.id))):
        if snapshot.status == NotionMatchStatus.linked and snapshot.student_id:
            snapshots_by_student[snapshot.student_id] = snapshot
    people: dict = defaultdict(lambda: defaultdict(list))
    for assignment in assignments:
        if assignment.is_active and assignment.mentor_id and assignment.student_id:
            person = user_by_id.get(assignment.mentor_id)
            if person:
                people[assignment.student_id][assignment.role].append(person.name)
    active_chat_ids = {
        session.student_id for session, chat in sessions
        if session.student_id and session.status == TelegramSessionStatus.active
        and chat.status == TelegramChatStatus.active
    }

    rows = []
    for student in students:
        if student_ids is not None and student.id not in student_ids:
            continue
        snapshot = snapshots_by_student.get(student.id)
        raw_status, normalized_status = notion_status(snapshot)
        contract = contracts_by_student.get(student.id)
        crm_status = contract.pipeline_status.value if contract and contract.pipeline_status else None
        access = bool(student.user_id and (user := user_by_id.get(student.user_id)) and user.is_active)
        mismatch = bool(snapshot and normalized_status and crm_status and normalized_status != crm_status)
        if dataset == "notion_active" and normalized_status != "active_work":
            continue
        if dataset == "platform_access" and not access:
            continue
        if dataset == "status_mismatch" and not mismatch:
            continue
        if only_active_notion and normalized_status != "active_work":
            continue
        if only_status_mismatch and not mismatch:
            continue
        if require_portal_access and not access:
            continue
        if portal_access == "yes" and not access or portal_access == "no" and access:
            continue
        role_people = people[student.id]
        legacy_mzk_id = getattr(contract, "mzk_manager_id", None) if contract else None
        if not role_people.get(MentorRole.mzk) and legacy_mzk_id in user_by_id:
            role_people[MentorRole.mzk].append(user_by_id[legacy_mzk_id].name)
        mentors = sorted({name for role, names in role_people.items() if role != MentorRole.mzk for name in names})
        if require_chat and student.id not in active_chat_ids:
            continue
        if active_chat == "yes" and student.id not in active_chat_ids or active_chat == "no" and student.id in active_chat_ids:
            continue
        if require_mzk and not role_people.get(MentorRole.mzk):
            continue
        if has_mzk == "yes" and not role_people.get(MentorRole.mzk) or has_mzk == "no" and role_people.get(MentorRole.mzk):
            continue
        if require_mentor and not mentors:
            continue
        if has_mentor == "yes" and not mentors or has_mentor == "no" and mentors:
            continue
        if notion_link == "yes" and not snapshot or notion_link == "no" and snapshot:
            continue
        portal_user = user_by_id.get(student.user_id) if student.user_id else None
        rows.append({
            "full_name": student.full_name or "",
            "phone": student.phone or "",
            "intake_year": str(student.intake_year or ""),
            "degree_level": student.degree_level.value if student.degree_level else "",
            "city": student.city or "",
            "email": getattr(portal_user, "email", "") if portal_user else "",
            "countries": ", ".join(sorted(countries_by_student[student.id])),
            "programs": ", ".join(sorted(programs_by_student[student.id])),
            "lead_mentor": ", ".join(sorted(leads_by_student[student.id])),
            "notion_status": raw_status,
            "crm_status": crm_status or "",
            "status_match": "Да" if snapshot and normalized_status and crm_status and not mismatch else "Нет" if mismatch else "Не сверено",
            "notion_link": "Связан" if snapshot else "Не связан",
            "notion_url": getattr(snapshot, "notion_url", "") or "" if snapshot else "",
            "portal_access": "Да" if access else "Нет",
            "active_chat": "Да" if student.id in active_chat_ids else "Нет",
            "mzk": ", ".join(sorted(set(role_people.get(MentorRole.mzk, [])))),
            "mentors": ", ".join(mentors),
        })
    # The Notion Overview can display rows that have not been linked to a CRM
    # student yet. Include exactly those explicitly passed by its current-view
    # export context; they never inflate the all-students/base export.
    if notion_snapshot_ids is not None:
        for snapshot in snapshots:
            if snapshot.id not in notion_snapshot_ids or snapshot.status != NotionMatchStatus.new or snapshot.student_id:
                continue
            raw_status, normalized_status = notion_status(snapshot)
            if dataset == "notion_active" and normalized_status != "active_work":
                continue
            if dataset == "status_mismatch":
                continue
            if only_status_mismatch:
                continue
            if portal_access == "yes" or active_chat == "yes" or has_mzk == "yes" or has_mentor == "yes" or notion_link == "yes":
                continue
            if notion_link == "no" and snapshot.status != NotionMatchStatus.new:
                continue
            rows.append({
                "full_name": snapshot.full_name or "", "phone": snapshot.phone_normalized or "",
                "intake_year": str((snapshot.normalized_data or {}).get("intake_raw") or ""),
                "degree_level": str((snapshot.normalized_data or {}).get("degree_raw") or ""),
                "city": "", "email": "", "countries": "", "programs": "", "lead_mentor": "",
                "notion_status": raw_status, "crm_status": "", "status_match": "Не сверено",
                "notion_link": "Не связан", "notion_url": snapshot.notion_url or "",
                "portal_access": "Нет", "active_chat": "Нет", "mzk": "", "mentors": "",
            })
    for field, options in (field_filters or {}).items():
        if options:
            selected = {option.casefold() for option in options}
            rows = [row for row in rows if selected.intersection(value.strip().casefold() for value in row.get(field, "").split(",") if value.strip())]
    if search.strip():
        needle = search.strip().casefold()
        rows = [row for row in rows if needle in " ".join(row.get(field, "") for field in ("full_name", "phone", "city", "email", "countries")).casefold()]
    return rows


def safe_cell(value: str) -> str:
    """Prefix formula-like CSV/TSV values so spreadsheet import treats them as text."""
    return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value


def safe_xlsx_cell(value: str) -> str:
    """XLSX stores strings with an explicit text type; only '=' is parsed as formula."""
    return "'" + value if value.startswith("=") else value


def serialize_rows(
    rows: list[dict[str, str]], columns: list[str], file_format: str,
    delimiter: str = ";", bom: bool = True, total_row: bool = False,
    parameters: dict[str, str] | None = None,
    split_by: str = "none",
) -> bytes:
    if file_format in {"csv", "tsv"}:
        if file_format == "tsv":
            delimiter, bom = "\t", False
        output = io.StringIO(newline="")
        writer = csv.writer(output, delimiter=delimiter)
        writer.writerow([COLUMNS[key] for key in columns])
        writer.writerows([safe_cell(row[key]) for key in columns] for row in rows)
        return (b"\xef\xbb\xbf" if bom else b"") + output.getvalue().encode("utf-8")
    if split_by not in {"none", "intake_year", "notion_status", "mzk"}:
        raise ValueError("Неверный способ разбиения на листы")
    workbook = Workbook()
    if split_by == "none":
        sheets = [("Студенты", rows)]
    else:
        split_labels = {"intake_year": "год набора", "notion_status": "статус", "mzk": "МЗК"}
        grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
        for row in rows:
            grouped[row.get(split_by, "").strip() or f"Без {split_labels[split_by]}"].append(row)
        sheets = sorted(grouped.items(), key=lambda item: (item[0].startswith("Без "), item[0].casefold())) or [("Студенты", [])]
    used_names: set[str] = set()
    for index, (name, sheet_rows) in enumerate(sheets):
        safe_name = "".join("_" if char in '[]:*?/\\' else char for char in name).strip("'")[:31] or f"Лист {index + 1}"
        original = safe_name
        suffix = 2
        while safe_name in used_names:
            tail = f" ({suffix})"
            safe_name = f"{original[:31 - len(tail)]}{tail}"
            suffix += 1
        used_names.add(safe_name)
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = safe_name
        _write_student_sheet(sheet, sheet_rows, columns, total_row)
    params_sheet = workbook.create_sheet("Параметры")
    params_sheet.append(["Параметр", "Значение"])
    from openpyxl.styles import Font, PatternFill
    header_fill = PatternFill("solid", fgColor="EEECE4")
    for cell in params_sheet[1]:
        cell.font = Font(bold=True, color="17170F")
        cell.fill = header_fill
    for key, value in (parameters or {}).items():
        params_sheet.append([safe_xlsx_cell(str(key)), safe_xlsx_cell(str(value))])
    params_sheet.column_dimensions["A"].width = 28
    params_sheet.column_dimensions["B"].width = 72
    params_sheet.freeze_panes = "A2"
    params_sheet.sheet_view.showGridLines = False
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_student_sheet(sheet, rows: list[dict[str, str]], columns: list[str], total_row: bool) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill("solid", fgColor="EEECE4")
    sheet.append([COLUMNS[key] for key in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17170F")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        sheet.append([safe_xlsx_cell(row[key]) for key in columns])
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if "Статус Notion" in columns:
            status_cell = sheet.cell(sheet.max_row, columns.index("Статус Notion") + 1)
            status = str(status_cell.value or "").lower()
            color = "DCEFE0" if "активная работа" in status else "FBE2CB" if "виз" in status or "ожид" in status else "F7D8D7" if "возврат" in status or "пропал" in status else "D6EAEE" if "поступ" in status else None
            if color:
                status_cell.fill = PatternFill("solid", fgColor=color)
    if total_row:
        sheet.append([f"Итого: {len(rows)} студентов", *[""] * max(0, len(columns) - 1)])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(48, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
