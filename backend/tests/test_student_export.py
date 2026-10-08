import csv
import io
import uuid
from openpyxl import load_workbook
from datetime import datetime, timezone
from types import SimpleNamespace as Obj

from app.models.contract import PipelineStatus
from app.models.mentor_assignment import MentorRole
from app.models.notion_snapshot import NotionMatchStatus
from app.models.student import DegreeLevel
from app.models.telegram_chat import TelegramChatStatus
from app.models.telegram_chat_session import TelegramSessionStatus
from app.services.student_export import make_rows, safe_cell, serialize_rows


def test_export_separates_notion_activity_from_platform_access_and_assignments():
    student_id, user_id, mentor_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    student = Obj(id=student_id, full_name="Алия", phone="123", intake_year=2027,
                  degree_level=DegreeLevel.undergraduate, city="Алматы", user_id=user_id)
    contract = Obj(id=uuid.uuid4(), student_id=student_id, created_at=datetime.now(timezone.utc),
                   pipeline_status=PipelineStatus.paused)
    snapshot = Obj(id=uuid.uuid4(), student_id=student_id, status=NotionMatchStatus.linked,
                   synced_at=datetime.now(timezone.utc), normalized_data={"payment_status_raw": "Активная работа"})
    users = [Obj(id=user_id, is_active=False, name="Алия"), Obj(id=mentor_id, is_active=True, name="Ментор")]
    assignment = Obj(student_id=student_id, mentor_id=mentor_id, is_active=True, role=MentorRole.lead)
    session = Obj(student_id=student_id, status=TelegramSessionStatus.active)
    chat = Obj(status=TelegramChatStatus.active)
    args = ([student], [contract], [snapshot], users, [assignment], [(session, chat)])

    row = make_rows(*args, "notion_active")[0]
    assert row["notion_status"] == "Активная работа"
    assert row["crm_status"] == "paused"
    assert row["status_match"] == "Нет"
    assert row["portal_access"] == "Нет"
    assert row["active_chat"] == "Да"
    assert row["mentors"] == "Ментор"
    assert make_rows(*args, "platform_access") == []
    assert len(make_rows(*args, "all", portal_access="no")) == 1
    assert make_rows(*args, "all", portal_access="yes") == []
    assert len(make_rows(*args, "all", active_chat="yes")) == 1
    assert make_rows(*args, "all", active_chat="no") == []
    assert len(make_rows(*args, "all", notion_link="yes")) == 1
    assert make_rows(*args, "all", notion_link="no") == []
    assert make_rows(*args, "all", require_portal_access=True, require_chat=True) == []
    assert len(make_rows(*args, "all", require_chat=True, require_mentor=True)) == 1
    assert make_rows(*args, "all", require_mzk=True) == []
    assert len(make_rows(*args, "status_mismatch")) == 1
    assert len(make_rows(*args, "all", only_active_notion=True, only_status_mismatch=True)) == 1
    assert make_rows(*args, "all", only_active_notion=True, only_status_mismatch=True, portal_access="yes") == []
    assert len(make_rows(*args, "all", search="али")) == 1
    assert make_rows(*args, "all", search="Арман") == []


def test_csv_uses_selected_columns_and_escapes_formulas():
    content = serialize_rows([{"full_name": "=1+1", "phone": "123"}], ["phone", "full_name"], "csv", delimiter=";")
    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert rows == [["Телефон", "Студент"], ["123", "'=1+1"]]
    assert safe_cell("\t=1+1") == "'\t=1+1"


def test_xlsx_has_parameter_sheet_and_total_row():
    content = serialize_rows(
        [{"full_name": "Алия", "phone": "+00123"}], ["full_name", "phone"], "xlsx",
        total_row=True, parameters={"Кто выгрузил": "Администратор", "Строк": "1"},
    )
    workbook = load_workbook(io.BytesIO(content), read_only=False)
    assert workbook.sheetnames == ["Студенты", "Параметры"]
    assert workbook["Студенты"]["A2"].value == "Алия"
    assert workbook["Студенты"]["B2"].value == "+00123"
    assert workbook["Студенты"]["A3"].value == "Итого: 1 студентов"
    assert workbook["Параметры"]["B2"].value == "Администратор"


def test_xlsx_splits_sheets_by_intake_and_keeps_parameter_sheet():
    content = serialize_rows(
        [
            {"full_name": "Алия", "intake_year": "2027"},
            {"full_name": "Арман", "intake_year": "2026"},
            {"full_name": "Сабина", "intake_year": "2027"},
        ],
        ["full_name", "intake_year"], "xlsx", split_by="intake_year",
        parameters={"Строк": "3"},
    )
    workbook = load_workbook(io.BytesIO(content), read_only=False)
    assert workbook.sheetnames == ["2026", "2027", "Параметры"]
    assert workbook["2026"]["A2"].value == "Арман"
    assert workbook["2027"]["A2"].value == "Алия"
    assert workbook["2027"]["A3"].value == "Сабина"
